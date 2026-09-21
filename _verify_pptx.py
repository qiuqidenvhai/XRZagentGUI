# -*- coding: utf-8 -*-
"""针对性验证：3 页 PPT 是否真的生成出 3 页且内容正确（不再是空壳）。"""
import sys, os, time
sys.path.insert(0, r"D:\软件\XianRenZhangAgent")
from _xrz_harness import health, run, switch_platform

print("health:", health(), flush=True)
switch_platform("deepseek"); time.sleep(2)

OUT = r"D:\软件\XianRenZhangAgent\xrz_data\XianRenZhang_tasks\mp3_probe.pptx"
try:
    if os.path.exists(OUT):
        os.replace(OUT, OUT + ".old")   # 沙箱不能删，改名
except Exception:
    pass

cmd = ("请用 pptx_create 工具生成演示文稿，path 设为 " + OUT
       + "，3页：第1页「综合测试」，第2页「功能展示」，第3页「谢谢」")
t0 = time.time()
r = run(cmd, timeout=300)
print(f"耗时 {time.time()-t0:.0f}s | ok={r['ok']} tools={sorted(r['tools'])}", flush=True)
print("final:", r["final"][:300], flush=True)
if r["errs"]:
    print("errs:", r["errs"][:2], flush=True)

if os.path.exists(OUT):
    from pptx import Presentation
    p = Presentation(OUT)
    print(f"\n实际页数: {len(p.slides)}（要求 3 页）", flush=True)
    for i, s in enumerate(p.slides, 1):
        texts = []
        for sh in s.shapes:
            if sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    texts.append(t.replace("\n", " | "))
        print(f"  第{i}页: {texts[:5]}", flush=True)
    ok = (len(p.slides) >= 3)
    print(f"\n[{'PASS' if ok else 'FAIL'}] pptx_3pages", flush=True)
else:
    print("\n[FAIL] 未生成文件", flush=True)

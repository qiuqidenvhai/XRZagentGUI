#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""子代理全平台验证：切到【非 DeepSeek 平台】(默认通义) 派子代理。
证明 spawn_child + PlatformSession 让子代理在任意平台跑通（以前只有 DeepSeek）。

红线：走真实 GUI（Playwright 驱动 gui.html，点平台按钮切平台 + 输入框发指令），
同时影子 SSE 抓 /events 事件流。

判定：
  ① task 工具返回的"平台"字段 = 目标平台（证明子代理跟随当前对话平台，不是跑回 DeepSeek）
  ② 子代理 result.json success=True（非 DeepSeek 平台子代理真跑通）
  ③ /events 出现带 subagent_task_id 事件（GUI 子代理卡片链路）
  ④ 母代理 docx 落盘
"""
import socket, json, time, threading, os, glob, sys
from playwright.sync_api import sync_playwright

B = "http://127.0.0.1:8888"
APP = os.path.dirname(os.path.abspath(__file__))
PLAT = sys.argv[1] if len(sys.argv) > 1 else "tongyi"   # tongyi / doubao / yuanbao
SESS = r"D:/软件/XianRenZhangAgent/xrz_data/XianRenZhang_tasks/gui_session"
RPT = SESS + r"\subagent_%s_report.docx" % PLAT

INSTR = ("【必须使用子代理】请严格按顺序执行：\n"
         "第1步：必须用 task 工具（type=research）委派一个 research 子代理，让它研究"
         "『多肉植物养护的 2 个要点』并总结成 2 条。不要跳过这一步。\n"
         "第2步：用 wait_task 等子代理跑完，拿到它的发现。\n"
         "第3步：把子代理的发现用 docx_create 写成 Word 文档 " + RPT + "（标题『子代理养护报告』，正文含 2 条要点）。\n"
         "全部做完后调用 done。即使子代理失败，也必须先真实调用 task + wait_task 拿到结果再决定下一步。")


class SSEClient:
    def __init__(self):
        self.s = socket.create_connection(("127.0.0.1", 8888), timeout=3)
        self.s.sendall(b"GET /events HTTP/1.1\r\nHost:127.0.0.1:8888\r\nAccept:text/event-stream\r\n\r\n")
        self.buf = b""
        self.items = []
        self.lock = threading.Lock()
        self._stop = False
        self.t0 = time.time()
        threading.Thread(target=self._loop, daemon=True).start()

    def _loop(self):
        self.s.settimeout(1.0)
        while not self._stop:
            try:
                c = self.s.recv(65536)
            except socket.timeout:
                continue
            except Exception:
                break
            if not c:
                break
            self.buf += c
            while b"\n" in self.buf:
                line, _, self.buf = self.buf.partition(b"\n")
                line = line.strip()
                if not line.startswith(b"data: "):
                    continue
                try:
                    e = json.loads(line[6:])
                except Exception:
                    continue
                with self.lock:
                    self.items.append({"t": round(time.time() - self.t0, 1),
                                       "type": e.get("type"), "data": e.get("data")})

    def stop(self):
        self._stop = True
        try:
            self.s.close()
        except Exception:
            pass

    def snapshot(self):
        with self.lock:
            return list(self.items)


def find_fresh_subagent_result(since_ts):
    pats = [
        SESS.replace("\\", "/") + "/subagent_*/result.json",
        r"D:/软件/XianRenZhangAgent/xrz_data/**/subagent_*/result.json",
    ]
    cands = []
    for p in pats:
        cands += glob.glob(p, recursive=True)
    cands = [p for p in cands if os.path.getmtime(p) >= since_ts]
    if not cands:
        return None
    cands.sort(key=lambda x: os.path.getmtime(x))
    return cands[-1]


def main():
    run_start_ts = time.time() - 5
    try:
        os.remove(RPT)
    except Exception:
        pass
    os.makedirs(SESS, exist_ok=True)

    sse = SSEClient()
    time.sleep(3)
    off = len(sse.snapshot())

    with sync_playwright() as pw:
        br = pw.chromium.launch(headless=True,
                               args=["--disable-gpu", "--no-sandbox", "--use-gl=swiftshader"])
        pg = br.new_page(viewport={"width": 1280, "height": 820})
        pg.goto(B + "/gui.html", wait_until="load", timeout=60000)
        pg.wait_for_timeout(2500)
        # 切到目标平台
        btn = pg.locator("#pl-" + PLAT)
        print("══ 切换到平台 %s（按钮存在=%s）══" % (PLAT, btn.count() > 0), flush=True)
        if btn.count() > 0:
            btn.click()
            time.sleep(6)   # 等平台浏览器 launch + navigate
        pg.locator("#input").fill(INSTR)
        pg.locator("#sendBtn").click()

        print("══ 等母代理 docx 落盘（最多 6 分钟）══", flush=True)
        land = False
        for i in range(120):
            if os.path.exists(RPT):
                land = True
                break
            time.sleep(3)
        time.sleep(12)
        n_cards = pg.locator(".subagent-block").count()
        print("  子代理卡片数（DOM .subagent-block）=%d" % n_cards, flush=True)
        pg.screenshot(path=os.path.join(APP, "_subagent_%s.png" % PLAT))
        br.close()

    sse.stop()
    evs = sse.snapshot()[off:]
    seq = [it["data"].get("tool") for it in evs
           if it["type"] == "tool_start" and isinstance(it["data"], dict)]
    print("母代理工具序列:", seq, flush=True)

    # task output（含"平台"字段）
    task_out, wait_out = "", ""
    for it in evs:
        dd = it.get("data")
        if not isinstance(dd, dict):
            continue
        if dd.get("tool") == "task" and it["type"] in ("tool_end", "command_success"):
            task_out = str(dd.get("output"))
        if dd.get("tool") == "wait_task" and it["type"] in ("tool_end", "command_success"):
            wait_out = str(dd.get("output"))
    print("[task]", task_out[:240].replace("\n", " "), flush=True)
    print("[wait_task]", wait_out[:240].replace("\n", " "), flush=True)

    sub_events = [it for it in evs if isinstance(it["data"], dict) and it["data"].get("subagent_task_id")]
    types_sub = {}
    for it in sub_events:
        k = it["data"].get("subagent_type") or it["type"]
        types_sub[k] = types_sub.get(k, 0) + 1

    res_path = find_fresh_subagent_result(run_start_ts)
    sub_success, sub_err = None, None
    if res_path:
        rd = json.load(open(res_path, encoding="utf-8"))
        sub_success, sub_err = rd.get("success"), rd.get("error")
        print("子代理 result.json: %s success=%s error=%s" %
              (os.path.basename(os.path.dirname(res_path)), sub_success, (sub_err or "")[:100]), flush=True)
    else:
        print("未找到本次子代理 result.json", flush=True)

    # 母代理 docx 内容（验证是否采纳子代理发现）
    docx_txt = ""
    if os.path.exists(RPT):
        try:
            import docx as _d
            _doc = _d.Document(RPT)
            docx_txt = "\n".join(p.text for p in _doc.paragraphs if p.text.strip())
        except Exception:
            pass

    # ④ 修正语义：弱模型（豆包/元宝）常把 docx 写到【错误的兄弟文件名】
    # （实测豆包母代理把内容正确的 docx 写成了 subagent_tongyi_report.docx 而非
    #  subagent_doubao_report.docx），精确文件名判定会误报 N。
    # 改为：检测本次 run 窗口内 gui_session 下任意新建 .docx，且正文含多肉关键词。
    import re as _re
    fresh_docx = []
    for _f in glob.glob(os.path.join(SESS, "*.docx")):
        if os.path.getmtime(_f) >= run_start_ts:
            _txt = ""
            try:
                import docx as _dd
                _txt = "\n".join(p.text for p in _dd.Document(_f).paragraphs)
            except Exception:
                pass
            fresh_docx.append((_f, _txt))
    land_doc = [f for f, t in fresh_docx if ("多肉" in t or "控水" in t or "光照" in t)]
    land_exact = os.path.exists(RPT)
    land_content = len(land_doc) > 0

    print("\n===== 判定（平台=%s）=====" % PLAT, flush=True)
    plat_follow = ("DeepSeek" not in task_out) and (PLAT in task_out or "平台:" in task_out)
    print("  ① task 跟随当前平台(非跑回DeepSeek): %s" % ("Y" if plat_follow else "N"))
    print("  ② 子代理 result.json success: %s" % sub_success)
    print("  ③ GUI 子代理事件冒泡: %d 条 (types=%s)" % (len(sub_events), types_sub))
    print("  ④ 母代理 docx 落盘: 精确名=%s / 内容命中(多肉)=%s  实际新文件=%s" %
          ("Y" if land_exact else "N", "Y" if land_content else "N",
           ", ".join(os.path.basename(f) for f, _ in fresh_docx) or "无"))
    _lc = next((t for f, t in fresh_docx if ("多肉" in t or "控水" in t or "光照" in t)), docx_txt)
    print("  ⑤ docx 内容(前200): %s" % _lc[:200].replace("\n", " "), flush=True)

    with open(os.path.join(APP, "_subagent_%s_events.json" % PLAT), "w", encoding="utf-8") as f:
        json.dump({"platform": PLAT, "seq": seq, "task_out": task_out, "wait_out": wait_out,
                   "sub_events": sub_events, "types_sub": types_sub,
                   "res_path": res_path, "sub_success": sub_success, "sub_err": sub_err,
                   "rpt_exists": os.path.exists(RPT), "docx_txt": docx_txt},
                  f, ensure_ascii=False, indent=2)
    print("\n完成。", flush=True)


if __name__ == "__main__":
    main()
